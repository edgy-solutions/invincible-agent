"""Generate the SAMPLE CDRL template: policy/overlays/sample/fin_export_profiles/cdrl_a001_ev_summary.xlsx.

Proposal `docs/proposals/a-measure-is-an-overlay-row-and-a-cdrl-is-an-export-profile.md`, Part 2.
The template is a deliverable's LAYOUT only (title, labels, a header row, column widths); every
number in a CDRL is written into it by `agent_fleet/finance_agent/cdrl.py` from a verb's answer.

The bytes are DETERMINISTIC, because the profile row pins them by sha256 (a template whose bytes
differ refuses the profile at load). openpyxl stamps the workbook properties and every zip member
with the current time, so both are fixed here: the properties to a constant, and the zip rewritten
with a fixed member timestamp and a fixed member order.

Run: uv run --no-sync python scripts/generate_cdrl_sample_template.py [--out PATH]
Prints the file's sha256.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[1]
DEFAULT_OUT = (REPO / "policy" / "overlays" / "sample" / "fin_export_profiles"
               / "cdrl_a001_ev_summary.xlsx")

SHEET = "EV Summary"
TITLE = "CDRL A001 - Notional Earned Value Summary (SYNTHETIC)"
LABELS = ("Program", "Program id", "As of", "CDRL item")   # A2:A5
HEADERS = ("Period", "Cum BCWS", "Cum BCWP", "Cum ACWP", "Cum CPI", "Cum SPI", "TCPI")  # row 7
WIDTHS = (34, 16, 16, 16, 12, 12, 12)

#: Constant workbook properties and zip member time: the file is a function of this script alone.
_FIXED_WHEN = datetime(2026, 10, 28, 0, 0, 0)
_FIXED_ZIP_TIME = (2026, 10, 28, 0, 0, 0)


def build_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws["A1"] = TITLE
    ws["A1"].font = Font(bold=True, size=14)
    for offset, label in enumerate(LABELS):
        ws.cell(row=2 + offset, column=1, value=label).font = Font(bold=True)
    for col, header in enumerate(HEADERS, start=1):
        ws.cell(row=7, column=col, value=header).font = Font(bold=True)
    for col, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    wb.properties.creator = "invincible-agent"
    wb.properties.lastModifiedBy = "invincible-agent"
    wb.properties.created = _FIXED_WHEN
    wb.properties.modified = _FIXED_WHEN

    raw = io.BytesIO()
    wb.save(raw)
    out = io.BytesIO()
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist()):
            info = zipfile.ZipInfo(name, date_time=_FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            data = src.read(name)
            if name == "docProps/core.xml":
                # openpyxl stamps `modified` with the wall clock at save, whatever the property says.
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-10-28T00:00:00Z", data)
            dst.writestr(info, data)
    return out.getvalue()


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()
    data = build_bytes()
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_bytes(data)
    print(hashlib.sha256(data).hexdigest())


if __name__ == "__main__":
    main()
