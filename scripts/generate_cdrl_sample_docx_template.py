"""Generate the SAMPLE docx CDRL template: policy/overlays/sample/fin_export_profiles/cdrl_a002_cost_performance.docx.

The docx twin of `scripts/generate_cdrl_sample_template.py`: a deliverable's LAYOUT only (a heading,
one sentence, a 5x2 title table with its labels, a 2x5 table, a header row and one EMPTY body row); every number in a CDRL is
written into it by `agent_fleet/finance_agent/cdrl.py` from a verb's answer.

The bytes are DETERMINISTIC, because the profile row pins them by sha256 (a template whose bytes
differ refuses the profile at load). python-docx stamps the core properties, so they are fixed to
the xlsx generator's constant, and the zip is rewritten with the xlsx generator's `rewrite_zip`
(fixed member timestamp, fixed member order).

Run: uv run --no-sync python scripts/generate_cdrl_sample_docx_template.py [--out PATH]
Prints the file's sha256.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import sys
from pathlib import Path

from docx import Document

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from generate_cdrl_sample_template import _FIXED_WHEN, rewrite_zip  # noqa: E402

DEFAULT_OUT = (REPO / "policy" / "overlays" / "sample" / "fin_export_profiles"
               / "cdrl_a002_cost_performance.docx")

TITLE = "CDRL A002 - Notional Cost Performance Report (SYNTHETIC)"
NOTE = "The data item description DI-NOTIONAL-0002 is notional; it names no real data item."
LABELS = ("Program", "Program id", "As of", "CDRL item", "Data item")   # table 1, column A
HEADERS = ("Period", "Cum BCWP", "Cum ACWP", "Cum CPI", "TCPI")        # table 2, row 1
BODY_ROWS = 1                                                           # table 2: one empty body row


def build_bytes() -> bytes:
    doc = Document()
    doc.add_heading(TITLE, level=1)
    doc.add_paragraph(NOTE)
    t1 = doc.add_table(rows=len(LABELS), cols=2)
    for i, label in enumerate(LABELS):
        t1.cell(i, 0).text = label
    t2 = doc.add_table(rows=1 + BODY_ROWS, cols=len(HEADERS))   # row 2 is EMPTY: the loader needs the start cell in the grid, and placeholder text would survive an absent value
    for j, header in enumerate(HEADERS):
        t2.cell(0, j).text = header
    cp = doc.core_properties
    cp.author = "invincible-agent"
    cp.last_modified_by = "invincible-agent"
    cp.created = _FIXED_WHEN
    cp.modified = _FIXED_WHEN
    cp.revision = 1
    raw = io.BytesIO()
    doc.save(raw)
    return rewrite_zip(raw.getvalue())


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()
    data = build_bytes()
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_bytes(data)
    print(hashlib.sha256(data).hexdigest())


if __name__ == "__main__":
    main()
