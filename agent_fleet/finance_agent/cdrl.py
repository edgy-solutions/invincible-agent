"""Produce one CDRL: an export profile's file (xlsx, docx, csv or json), filled with values verbs returned.

Proposal `docs/proposals/a-measure-is-an-overlay-row-and-a-cdrl-is-an-export-profile.md`, Part 2.

THE PRODUCER COMPUTES NOTHING. Each source is run through the engine's one envelope builder (the
function `/measure/{fn}` serves, handed in as `run_source` so this module never imports `main`),
the rows are joined on `table.key`, and every number written is the float the envelope carries.
A period a source marks `undefined` writes `undefined: <reason>` into the cell, never a number.

THE CLAIM IS `cells_sha256`, NOT THE FILE'S BYTES. An xlsx is a zip with write times in it, so the
file is not reproducible; the hash of every cell this module WROTE (container, coordinate, value) in
canonical JSON is. The xlsx `_manifest` sheet carries the envelope, which carries that hash, so the
manifest is not itself one of the hashed cells.

FOUR WRITERS, ONE CONTRACT. `build_cdrl` runs the sources and joins the keys ONCE, then dispatches on
`profile.format` to `_write_xlsx`, `_write_docx`, `_write_csv` or `_write_json`. Each returns the
`[container, coord, value]` list it WROTE (value = what the FILE holds) and a `save(envelope, path)`.
Containers: the sheet title (xlsx), `T<n>` (docx), `csv`, `json`; coords: A1 (xlsx, csv), the table's
A1 (docx), a `/a/b` pointer (json).
"""
from __future__ import annotations

import hashlib
import json
import os
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Optional

try:  # flat in the image (/app), packaged in the repo -- see `measures.py`
    from export import MANIFEST_SCHEMA
    from export_profiles import TEXT_FORMATS, ExportProfile
except ImportError:
    from agent_fleet.finance_agent.export import MANIFEST_SCHEMA
    from agent_fleet.finance_agent.export_profiles import TEXT_FORMATS, ExportProfile

#: What `run_source` returns for one source: the served envelope and the method identity.
Source = tuple[dict[str, Any], dict[str, Any]]
RunSource = Callable[[str, dict[str, Any]], Source]

#: THE closed map from a delivered file's extension to its media type (`/artifact` reads it).
MEDIA_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "csv": "text/csv; charset=utf-8",
    "json": "application/json",
}

_METHODS_HEADER = ("fn", "measure_id", "name", "version", "module_sha256")


def _canonical(obj: Any) -> str:
    return json.dumps(obj, sort_keys=True, separators=(",", ":"), default=str)


def _sha(obj: Any) -> str:
    return hashlib.sha256(_canonical(obj).encode("utf-8")).hexdigest()


def artifact_filename(profile_name: str, program_id: str, fmt: str) -> str:
    """THE ONE PLACE the file is named: the engine's `/artifact` allow-list is derived from it.
    `fmt` is the profile's format and is the extension."""
    return f"fin-cdrl-{profile_name}-{program_id}.{fmt}"


def _norm(value: Any) -> Any:
    """What a file holds for a float: 16 significant digits (what openpyxl stores), an int when
    whole. The ONE normalisation every writer shares, so a reader's recomputation from the file
    matches `cells_sha256` (1 ulp from the envelope's float)."""
    if isinstance(value, float):
        value = float(f"{value:.16g}")
        if value.is_integer() and abs(value) < 1e15:
            value = int(value)   # the file stores `1000`, so a reader gets an int back
    return value


def _text(value: Any, number_format: Optional[str]) -> str:
    """A value as the TEXT a docx or csv cell holds: a number through `TEXT_FORMATS` when the
    column names a format, else `str()` of the normalised number; a string as it is."""
    value = _norm(value)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if number_format:
            return TEXT_FORMATS[number_format].format(value)
        return str(value)
    return str(value)


def _fact(profile: ExportProfile, name: str, program: Any, as_of: str) -> str:
    if name == "program.name":
        return program.name
    if name == "program.id":
        return program.program_id
    if name == "as_of":
        return as_of
    return profile.cdrl[name.split(".", 1)[1]]


def _cell_value(row: Optional[dict[str, Any]], field: str) -> Any:
    """What one table cell holds: the envelope's value, `undefined: <reason>` where the row says
    the period is undefined, else nothing (a period a source did not return is an empty cell)."""
    if row is None:
        return None
    value = row.get(field)
    if value is None and row.get("undefined"):
        return f"undefined: {row['undefined']}"
    return value


@dataclass
class _Run:
    """What every writer shares: the program, the joined keys in order, the sources' methods."""
    profile: ExportProfile
    program: Any
    as_of: str
    order: list[Any]
    keyed: dict[str, dict[Any, dict[str, Any]]]
    methods: dict[str, dict[str, Any]]

    def facts(self) -> list[tuple[str, Any]]:
        return [(c, _fact(self.profile, f, self.program, self.as_of))
                for c, f in self.profile.cells.items()]

    def table(self) -> list[list[Any]]:
        """One list per key, one value per column (None = nothing to write)."""
        return [
            [key if col.source == "" else _cell_value(self.keyed[col.source].get(key), col.field)
             for col in self.profile.columns]
            for key in self.order
        ]


Saver = Callable[[dict[str, Any], Path], None]


def _write_xlsx(run: _Run) -> tuple[list[list[Any]], Saver]:
    import openpyxl
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

    profile = run.profile
    wb = openpyxl.load_workbook(profile.template_path)
    ws = wb[profile.sheet]
    written: list[list[Any]] = []

    def put(sheet: Any, coord: str, value: Any, number_format: Optional[str] = None) -> None:
        if value == "":
            return   # a file has no empty-string cell; it would hash here and read back as nothing
        value = _norm(value)   # hash what the FILE holds
        sheet[coord] = value
        if number_format and isinstance(value, (int, float)) and not isinstance(value, bool):
            sheet[coord].number_format = number_format
        written.append([sheet.title, coord, value])

    for coord, value in run.facts():
        put(ws, coord, value)
    col_letter, first_row = coordinate_from_string(profile.table_start)
    first_col = column_index_from_string(col_letter)
    for i, values in enumerate(run.table()):
        for j, (col, value) in enumerate(zip(profile.columns, values)):
            coord = ws.cell(row=first_row + i, column=first_col + j).coordinate
            if value is not None:
                put(ws, coord, value, col.format)

    # ── the methods sheet: one row per source, the method that produced its figures ────────
    if profile.provenance_sheets:
        ms = wb.create_sheet(profile.methods_sheet)
        for j, head in enumerate(_METHODS_HEADER, start=1):
            put(ms, ms.cell(row=1, column=j).coordinate, head)
        for i, fn in enumerate(profile.sources, start=2):
            m = run.methods[fn]
            for j, value in enumerate(
                    (fn, m.get("measure_id", ""), m.get("name", ""), m.get("version", ""),
                     m.get("module_sha256", "")), start=1):
                put(ms, ms.cell(row=i, column=j).coordinate, value)

    def save(envelope: dict[str, Any], path: Path) -> None:
        # The manifest is written AFTER the hash and is not part of it. A profile with
        # `provenance_sheets: false` delivers exactly the template's sheets; the envelope (JSON)
        # and the BFF's X-Cells-Sha256 header carry the provenance instead.
        if profile.provenance_sheets:
            mf = wb.create_sheet(profile.manifest_sheet)
            mf["A1"] = _canonical(envelope)
            mf.sheet_state = "hidden"
        wb.save(path)

    return written, save


def _set_cell_text(cell: Any, text: str) -> None:
    """THE DOCX WRITE METHOD: replace the text of the cell's FIRST paragraph's FIRST run (blanking
    any further runs, so the template's run style survives), drop the cell's other paragraphs, and
    fall back to `cell.text = text` only when the first paragraph has no run. After it,
    `cell.text == text`, which is what the hash claims."""
    paragraphs = cell.paragraphs
    for extra in paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    runs = paragraphs[0].runs if paragraphs else []
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        cell.text = text


def _write_docx(run: _Run) -> tuple[list[list[Any]], Saver]:
    import docx
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

    profile = run.profile
    document = docx.Document(str(profile.template_path))
    tables = document.tables
    written: list[list[Any]] = []

    def put(n: int, coord: str, text: str) -> None:
        if text == "":
            return
        letters, row = coordinate_from_string(coord)
        _set_cell_text(tables[n - 1].cell(row - 1, column_index_from_string(letters) - 1), text)
        written.append([f"T{n}", coord, text])

    def split(tcell: str) -> tuple[int, str]:
        t, _, a1 = tcell.partition("!")
        return int(t[1:]), a1

    for tcell, value in run.facts():
        n, a1 = split(tcell)
        put(n, a1, _text(value, None))
    n, a1 = split(profile.table_start)
    col_letter, first_row = coordinate_from_string(a1)
    first_col = column_index_from_string(col_letter)
    table = tables[n - 1]
    data = run.table()
    while len(table.rows) < first_row - 1 + len(data):
        table.add_row()
    for i, values in enumerate(data):
        for j, (col, value) in enumerate(zip(profile.columns, values)):
            if value is not None:
                put(n, f"{get_column_letter(first_col + j)}{first_row + i}", _text(value, col.format))

    def save(_envelope: dict[str, Any], path: Path) -> None:
        document.save(str(path))

    return written, save


def _write_csv(run: _Run) -> tuple[list[list[Any]], Saver]:
    import csv
    from openpyxl.utils.cell import get_column_letter

    profile = run.profile
    grid: list[list[str]] = [[col.as_ for col in profile.columns]]
    for values in run.table():
        grid.append(["" if v is None else _text(v, col.format)
                     for col, v in zip(profile.columns, values)])
    written = [["csv", f"{get_column_letter(j)}{i}", text]
               for i, line in enumerate(grid, start=1)
               for j, text in enumerate(line, start=1) if text != ""]

    def save(_envelope: dict[str, Any], path: Path) -> None:
        with open(path, "w", encoding="utf-8", newline="") as fh:
            csv.writer(fh, lineterminator="\r\n").writerows(grid)

    return written, save


def _json_set(obj: dict[str, Any], pointer: str, value: Any) -> None:
    *parents, last = pointer[1:].split("/")
    for seg in parents:
        obj = obj.setdefault(seg, {})
    obj[last] = value


def _write_json(run: _Run) -> tuple[list[list[Any]], Saver]:
    profile = run.profile
    obj: dict[str, Any] = {}
    written: list[list[Any]] = []
    for pointer, value in run.facts():
        if value == "":
            continue
        value = _norm(value)
        _json_set(obj, pointer, value)
        written.append(["json", pointer, value])
    rows: list[dict[str, Any]] = []
    for i, values in enumerate(run.table()):
        row: dict[str, Any] = {}
        for col, value in zip(profile.columns, values):
            if value is None or value == "":
                continue   # an absent value is an omitted key, never null
            value = _norm(value)
            row[col.as_] = value
            written.append(["json", f"{profile.table_at}/{i}/{col.as_}", value])
        rows.append(row)
    _json_set(obj, profile.table_at, rows)

    def save(_envelope: dict[str, Any], path: Path) -> None:
        # Keys in DECLARATION order, never sorted: a DID's schema order is its form.
        path.write_bytes((json.dumps(obj, indent=2, ensure_ascii=False) + "\n").encode("utf-8"))

    return written, save


_WRITERS: dict[str, Callable[[_Run], tuple[list[list[Any]], Saver]]] = {
    "xlsx": _write_xlsx, "docx": _write_docx, "csv": _write_csv, "json": _write_json,
}


def build_cdrl(
    state: Any,
    profile: ExportProfile,
    *,
    program_id: str,
    as_of: str,
    out_dir: Path,
    run_source: RunSource,
) -> tuple[Path, dict[str, Any]]:
    """Write `fin-cdrl-<name>-<program_id>.<format>` into `out_dir`; return (path, envelope)."""
    params = {"program_id": program_id}
    served: dict[str, Source] = {fn: run_source(fn, dict(params)) for fn in profile.sources}
    rows_by_source = {fn: list(env.get("rows") or []) for fn, (env, _m) in served.items()}

    program = next((p for p in state.programs if p.program_id == program_id), None)
    if program is None:  # unreachable behind a source that refused it; named anyway
        raise ValueError(f"program {program_id!r} is not in the model")

    # ── the table: join on the key, in order of first appearance across the sources ────────
    keyed = {fn: {r[profile.table_key]: r for r in rows} for fn, rows in rows_by_source.items()}
    order: list[Any] = []
    for fn in profile.sources:
        for r in rows_by_source[fn]:
            if r[profile.table_key] not in order:
                order.append(r[profile.table_key])
    methods: dict[str, dict[str, Any]] = {fn: m for fn, (_e, m) in served.items()}

    written, save = _WRITERS[profile.format](
        _Run(profile, program, as_of, order, keyed, methods))

    cells_sha256 = _sha(sorted(written, key=lambda c: (c[0], c[1])))
    envelope = {
        "schema": MANIFEST_SCHEMA,
        "kind": "cdrl",
        "format": profile.format,
        "profile": profile.name,
        "template_sha256": profile.template_sha256,
        "program_id": program_id,
        "as_of": as_of,
        "sources": [
            {"fn": fn, "params": dict(params), "method": methods[fn],
             "rows_sha256": _sha(rows_by_source[fn])}
            for fn in profile.sources
        ],
        "cells_sha256": cells_sha256,
        "provenance_sheets": profile.provenance_sheets,
    }

    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / artifact_filename(profile.name, program_id, profile.format)
    fd, tmp = tempfile.mkstemp(dir=out_dir, suffix=f".{profile.format}.tmp")
    os.close(fd)
    try:
        save(envelope, Path(tmp))
        os.replace(tmp, path)   # a reader never sees half a file
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)
    return path, envelope
