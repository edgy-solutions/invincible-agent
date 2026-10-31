"""The CDRL docx path, proven with one synthetic data item: A002 / DI-NOTIONAL-0002.

Why: `cdrl._write_docx` and the docx branch of `export_profiles` were exercised only by a throwaway
profile built inside `test_a_cdrl_profile_names_its_format.py`. Here the SHIPPED sample overlay
carries a docx profile (`cdrl_a002_cost_performance_docx` over a committed, deterministic template),
so the path is proven on bytes a roll actually delivers. No BFF or route change: the engine's
`/cdrl_profiles` lists it, and the BFF relays that list.

Arms: 1 the template is deterministic and pinned (and the xlsx generator's output is unchanged),
2 the real loader declares the profile, 3 produced through engine-fin `/cdrl_export` as a member
(facts, the DID cell, the header row, one row per period, no stray template text, an absent value
is an empty cell), 4 the claim (reproducible, recomputed from the delivered file alone, served to a
member and refused to a non-member), 5 the BFF picker lists it, from a served engine-fin.

Fixtures are imported, not copied, from `test_a_cdrl_is_an_export_profile.py` (the engine client,
the program, ALICE / BOB) and `test_a_cdrl_profile_names_its_format.py` (the docx claim recompute,
which was inline there and is now a function: `docx_claim`).

Run: uv run --no-sync pytest tests/finance/test_a_docx_cdrl_from_a_synthetic_did.py -q
"""
from __future__ import annotations

import asyncio
import subprocess
import sys
from pathlib import Path
from types import SimpleNamespace

import docx
import yaml

REPO = Path(__file__).resolve().parents[2]
for _p in (REPO, REPO / "src", REPO / "scripts"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

from agent_fleet.finance_agent import cdrl, export_profiles  # noqa: E402
from agent_fleet.finance_agent import main as engine  # noqa: E402
from tests.finance.test_a_cdrl_is_an_export_profile import (  # noqa: E402,F401
    ALICE, BOB, PROGRAM, SAMPLE, ROW as A001_ROW, declared, profiles, served, _rows,
)
from tests.finance.test_a_cdrl_profile_names_its_format import (  # noqa: E402
    EXPECT, _Req, _User, _spent_state, docx_claim,
)

NAME = "cdrl_a002_cost_performance_docx"
ROW = SAMPLE / f"{NAME}.yaml"
TEMPLATE = SAMPLE / "cdrl_a002_cost_performance.docx"
AS_OF = "2026-10-28"
MEDIA = dict(EXPECT.values())["docx"]
IDX = "fin_performance_indices"
FILE = f"fin-cdrl-{NAME}-{PROGRAM}.docx"


def _raw() -> dict:
    return yaml.safe_load(ROW.read_text(encoding="utf-8"))


def _run_gen(script: str, out: Path) -> str:
    r = subprocess.run([sys.executable, str(REPO / "scripts" / script), "--out", str(out)],
                       capture_output=True, text=True, check=True)
    return r.stdout.strip()


def _post(c, user=ALICE):
    return c.post("/cdrl_export", json={
        "profile": NAME, "program_id": PROGRAM, "user_email": user, "as_of": AS_OF})


def _render(value, fmt):
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return export_profiles.TEXT_FORMATS[fmt].format(cdrl._norm(value))


def _expected_table(idx, tc):
    """[period, cum_bcwp, cum_acwp, cum_cpi, tcpi] as the TEXT the docx must hold, per period."""
    out = []
    for p in idx:
        out.append([p, _render(idx[p]["cum_bcwp"], "#,##0"), _render(idx[p]["cum_acwp"], "#,##0"),
                    _render(idx[p]["cum_cpi"], "0.000"),
                    _render(tc[p]["tcpi"] if p in tc else None, "0.000")])
    return out


def _cells(table):
    return [[c.text for c in row.cells] for row in table.rows]


# -- arm 1: the template is deterministic and pinned ---------------------------------------------

def test_1_the_generator_is_deterministic_equals_the_committed_template_and_the_pin(tmp_path):
    a, b = tmp_path / "a.docx", tmp_path / "b.docx"
    sha_a = _run_gen("generate_cdrl_sample_docx_template.py", a)
    sha_b = _run_gen("generate_cdrl_sample_docx_template.py", b)
    assert a.read_bytes() == b.read_bytes() == TEMPLATE.read_bytes()
    assert sha_a == sha_b == _raw()["template"]["sha256"]


def test_1_the_xlsx_generators_output_is_unchanged_by_the_shared_zip_rewrite(tmp_path):
    out = tmp_path / "a001.xlsx"
    sha = _run_gen("generate_cdrl_sample_template.py", out)
    pin = yaml.safe_load(A001_ROW.read_text(encoding="utf-8"))["template"]["sha256"]
    assert sha == pin
    assert out.read_bytes() == (SAMPLE / "cdrl_a001_ev_summary.xlsx").read_bytes()


# -- arm 2: it loads -------------------------------------------------------------------------------

def test_2_the_real_loader_declares_the_docx_profile(declared):
    got = {p.name: p for p in export_profiles.load_export_profiles(
        policy_dir=SAMPLE.parent / "no-seed", overlay_dirs=[SAMPLE], declared=declared)}
    p = got[NAME]
    assert p.format == "docx"
    assert p.cdrl["data_item"] == _raw()["cdrl"]["data_item"]


# -- arm 3: produced through engine-fin /cdrl_export as ALICE ------------------------------------

def test_3_the_delivered_file_holds_each_fact_the_header_row_and_one_row_per_period(served):
    c, dist, _ = served
    r = _post(c)
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["artifact_filename"] == FILE
    doc = docx.Document(str(dist / FILE))
    program = next(p for p in engine.STATE.programs if p.program_id == PROGRAM)
    raw = _raw()
    t1 = _cells(doc.tables[0])
    assert [row[1] for row in t1] == [program.name, PROGRAM, AS_OF, raw["cdrl"]["item"],
                                      raw["cdrl"]["data_item"]]
    assert [row[0] for row in t1] == ["Program", "Program id", "As of", "CDRL item", "Data item"]

    tpl = docx.Document(str(TEMPLATE))
    t2, tpl2 = _cells(doc.tables[1]), _cells(tpl.tables[1])
    assert t2[0] == tpl2[0]
    idx, tc = _rows(IDX), _rows("fin_tcpi")
    assert len(idx) >= 2
    want = _expected_table(idx, tc)
    assert t2[1:] == want
    assert len(t2) == 1 + len(idx)
    allowed = set(tpl2[0]) | {v for row in want for v in row} | {""}
    assert {v for row in t2 for v in row} <= allowed


def test_3_an_undefined_period_shows_the_reason_string(profiles, tmp_path, monkeypatch):
    state = _spent_state(monkeypatch)
    path, _ = cdrl.build_cdrl(
        state, profiles[NAME], program_id=PROGRAM, as_of=AS_OF, out_dir=tmp_path,
        run_source=lambda fn, params: engine._cdrl_source(fn, params))
    t2 = _cells(docx.Document(str(path)).tables[1])
    assert t2[1:] and {row[4] for row in t2[1:]} == {"undefined: budget_spent"}


def test_3_an_absent_value_is_an_empty_cell_and_no_template_text_survives(profiles, tmp_path):
    """The first period's TCPI is not returned: the body row the template ships must come back empty
    there, not carry any text of the template's."""
    def source(fn, params):
        env, method = engine._cdrl_source(fn, params)
        if fn == "fin_tcpi":
            env = {**env, "rows": list(env["rows"])[1:]}
        return env, method

    path, _ = cdrl.build_cdrl(engine.STATE, profiles[NAME], program_id=PROGRAM, as_of=AS_OF,
                              out_dir=tmp_path, run_source=source)
    t2 = _cells(docx.Document(str(path)).tables[1])
    tpl2 = _cells(docx.Document(str(TEMPLATE)).tables[1])
    idx = _rows(IDX)
    tc = {k: v for k, v in _rows("fin_tcpi").items() if k != next(iter(idx))}
    want = _expected_table(idx, tc)
    assert want[0][4] == "" and t2[1:] == want
    assert {v for row in t2 for v in row} <= set(tpl2[0]) | {v for row in want for v in row} | {""}


# -- arm 4: the claim ------------------------------------------------------------------------------

def test_4_the_claim_is_reproducible_recomputed_from_the_file_and_served_to_a_member_only(
        served, monkeypatch):
    c, dist, _ = served
    first = _post(c).json()
    claim = first["envelope"]["cells_sha256"]
    assert claim
    assert _post(c).json()["envelope"]["cells_sha256"] == claim
    assert docx_claim(dist / FILE, TEMPLATE) == claim

    with monkeypatch.context() as m:
        m.setattr(engine, "_current_caller", lambda: SimpleNamespace(authz_id=BOB))
        assert c.get(f"/artifact/{FILE}").status_code == 403
    with monkeypatch.context() as m:
        m.setattr(engine, "_current_caller", lambda: SimpleNamespace(authz_id=ALICE))
        r = c.get(f"/artifact/{FILE}")
    assert r.status_code == 200
    assert r.headers["content-type"] == MEDIA
    assert r.content == (dist / FILE).read_bytes()


# -- arm 5: it reaches the BFF picker with no BFF change -----------------------------------------

def test_5_the_bff_picker_lists_what_engine_fin_serves(served, monkeypatch):
    """The BFF's engine-fin client is the served engine-fin's own TestClient, so the list the BFF
    returns is what engine-fin answered."""
    import iagent.gateway as gw
    c, _, _ = served

    class _Client:
        def __init__(self, *a, **k):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *exc):
            return False

        async def get(self, url, headers=None):
            assert url.startswith(gw._ENGINE_FIN_URL)
            return c.get(url[len(gw._ENGINE_FIN_URL):], headers=headers)

    monkeypatch.setattr(gw.httpx, "AsyncClient", _Client)
    got = asyncio.run(gw.export_cdrl_profiles(_Req(), _User()))
    assert got == c.get("/cdrl_profiles").json()
    row = next(p for p in got if p["name"] == NAME)
    assert row["label"] == _raw()["label"] and row["format"] == "docx"
