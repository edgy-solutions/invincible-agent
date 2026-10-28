"""A CDRL is an export profile: one synthetic tabular CDRL, every cell a verb's envelope value.

Proposal `docs/proposals/a-measure-is-an-overlay-row-and-a-cdrl-is-an-export-profile.md`, Part 2.
Arms: 1 template determinism, 2 load refusals, 3 ROW_FIELDS cannot drift, 4 the file,
5 reproducibility, 6 undefined, 7 engine routes. (Arm 8, the BFF, is tests/test_bff_cdrl_export.py.)

Run: uv run --no-sync pytest tests/finance/test_a_cdrl_is_an_export_profile.py -q
"""
from __future__ import annotations

import copy
import hashlib
import json
import sys
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
import yaml
from fastapi.testclient import TestClient

REPO = Path(__file__).resolve().parents[2]
for _p in (REPO, REPO / "src", REPO / "scripts"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

from agent_fleet.finance_agent import cdrl, export_profiles, measures, seed  # noqa: E402
from agent_fleet.finance_agent import main as engine  # noqa: E402
import generate_cdrl_sample_template as gen  # noqa: E402

SAMPLE = REPO / "policy" / "overlays" / "sample" / "fin_export_profiles"
ROW = SAMPLE / "cdrl_a001_ev_summary.yaml"
TEMPLATE = SAMPLE / "cdrl_a001_ev_summary.xlsx"
PROGRAM = seed.PROGRAM_ID
ALICE, BOB = "alice@example.com", "bob@example.com"
PROFILE = "cdrl_a001_ev_summary"
ENVELOPE = {
    "error": "program_not_visible",
    "refusal": {"policy_id": "program_members",
                "missing_binding": "can_view_program(program)", "candidates": []},
    "message": "You are not a member of the program this request names.",
}


@pytest.fixture
def declared(monkeypatch):
    monkeypatch.setenv("MEASURE_VERB_OVERLAY_DIRS", str(REPO / "policy/overlays/sample/measure_verbs"))
    rows = engine._load_declared()
    monkeypatch.setattr(engine, "DECLARED", rows)
    monkeypatch.setitem(engine.VERB_TO_FN, "mesh:finToCompletePerformanceIndex", "fin_tcpi")
    return rows


@pytest.fixture
def profiles(declared, monkeypatch):
    monkeypatch.setenv("FIN_EXPORT_PROFILE_OVERLAY_DIRS", str(SAMPLE))
    p = {x.name: x for x in export_profiles.load_export_profiles(declared=declared)}
    monkeypatch.setattr(engine, "PROFILES", p)
    return p


def _row():
    return yaml.safe_load(ROW.read_text(encoding="utf-8"))


def _overlay(tmp_path, mutate, template_bytes=None):
    """A one-profile overlay dir: the sample row after `mutate(raw)`, plus its template
    (`template_bytes=False` writes none)."""
    raw = _row()
    mutate(raw)
    d = tmp_path / "ov"
    d.mkdir()
    if template_bytes is not False:
        (d / "cdrl_a001_ev_summary.xlsx").write_bytes(
            template_bytes if template_bytes is not None else TEMPLATE.read_bytes())
    (d / "row.yaml").write_text(yaml.safe_dump(raw), encoding="utf-8")
    return d


def _load(d, declared):
    return export_profiles.load_export_profiles(
        policy_dir=d.parent / "no-seed", overlay_dirs=[d], declared=declared)


# -- arm 1: the template is deterministic ------------------------------------------------------

def test_1_the_generator_is_deterministic_and_equals_the_committed_template():
    a, b = gen.build_bytes(), gen.build_bytes()
    assert a == b
    assert a == TEMPLATE.read_bytes()
    assert hashlib.sha256(a).hexdigest() == _row()["template"]["sha256"]


# -- arm 2: load refusals, one each, naming the file and the reason ----------------------------

def _refuses(tmp_path, declared, mutate, fragment, template_bytes=None):
    d = _overlay(tmp_path, mutate, template_bytes)
    with pytest.raises(export_profiles.ExportProfileError) as ei:
        _load(d, declared)
    msg = str(ei.value)
    assert "row.yaml" in msg, msg
    assert fragment in msg, msg


def test_2_the_sample_loads(declared, tmp_path):
    got = _load(_overlay(tmp_path, lambda r: None), declared)
    assert [p.name for p in got] == [PROFILE]


def test_2_template_sha_mismatch(declared, tmp_path):
    _refuses(tmp_path, declared, lambda r: None, "sha256", template_bytes=b"not the template")


def test_2_template_missing(declared, tmp_path):
    _refuses(tmp_path, declared, lambda r: None, "template", template_bytes=False)


def test_2_unknown_sheet(declared, tmp_path):
    _refuses(tmp_path, declared, lambda r: r.update(sheet="Nope"), "Nope")


def test_2_cell_fact_outside_the_vocabulary(declared, tmp_path):
    _refuses(tmp_path, declared, lambda r: r["cells"].update(B2="program.secret"), "program.secret")


def test_2_bad_coordinate(declared, tmp_path):
    _refuses(tmp_path, declared, lambda r: r["cells"].update({"ZZZ99999x": "as_of"}), "ZZZ99999x")


def test_2_unknown_source_fn(declared, tmp_path):
    _refuses(tmp_path, declared, lambda r: r["sources"].append({"fn": "fin_no_such_verb"}),
             "fin_no_such_verb")


def test_2_column_field_not_in_the_fns_field_set(declared, tmp_path):
    def m(r):
        r["table"]["columns"][1] = {"from": "fin_performance_indices.no_such_field"}
    _refuses(tmp_path, declared, m, "no_such_field")


def test_2_column_fn_not_in_sources(declared, tmp_path):
    def m(r):
        r["table"]["columns"][1] = {"from": "fin_burn_rate.burn"}
    _refuses(tmp_path, declared, m, "fin_burn_rate")


def test_2_a_formula_key(declared, tmp_path):
    def m(r):
        r["table"]["columns"][1] = {"from": "fin_performance_indices.cum_cpi", "formula": "=A1*2"}
    _refuses(tmp_path, declared, m, "formula")


def test_2_builtin_not_in_ROW_FIELDS(declared, tmp_path):
    builtin = next(f for f in measures.OUTPUT_URI if f not in measures.ROW_FIELDS)
    _refuses(tmp_path, declared, lambda r: r["sources"].append({"fn": builtin}), builtin)


# -- arm 3: ROW_FIELDS cannot drift ------------------------------------------------------------

@pytest.mark.parametrize("fn", sorted(measures.ROW_FIELDS))
def test_3_every_declared_field_is_carried_by_the_rows(fn):
    state = seed.build_seed()
    out = getattr(measures, fn)(state, program_id=PROGRAM)
    rows = out["rows"] if isinstance(out, dict) else out
    assert rows
    for r in rows:
        missing = measures.ROW_FIELDS[fn] - set(r)
        assert not missing, f"{fn}: rows lack declared fields {sorted(missing)}"


# -- arm 4: the file ---------------------------------------------------------------------------

def _build(profiles, out, state=None):
    return cdrl.build_cdrl(
        state or engine.STATE, profiles[PROFILE], program_id=PROGRAM, as_of="2026-10-28",
        out_dir=out, run_source=lambda fn, params: engine._cdrl_source(fn, params))


def _rows(fn):
    return {r["period"]: r for r in engine._measure_envelope(fn, {"program_id": PROGRAM})["rows"]}


def test_4_the_title_block(profiles, tmp_path):
    path, _ = _build(profiles, tmp_path)
    ws = openpyxl.load_workbook(path)["EV Summary"]
    program = next(p for p in engine.STATE.programs if p.program_id == PROGRAM)
    assert ws["B2"].value == program.name
    assert ws["B3"].value == PROGRAM
    assert ws["B4"].value == "2026-10-28"
    assert ws["B5"].value == "A001"


def test_4_one_row_per_period_in_order_each_cell_the_envelopes_value(profiles, tmp_path):
    path, _ = _build(profiles, tmp_path)
    ws = openpyxl.load_workbook(path)["EV Summary"]
    idx, tc = _rows("fin_performance_indices"), _rows("fin_tcpi")
    periods = list(idx)
    assert len(periods) >= 2
    assert [ws.cell(row=8 + i, column=1).value for i in range(len(periods))] == periods
    assert ws.cell(row=8 + len(periods), column=1).value is None, "an extra table row"
    spec = [("cum_bcws", idx, 2, "#,##0"), ("cum_bcwp", idx, 3, "#,##0"), ("cum_acwp", idx, 4, "#,##0"),
            ("cum_cpi", idx, 5, "0.000"), ("cum_spi", idx, 6, "0.000"), ("tcpi", tc, 7, "0.000")]
    for i, period in enumerate(periods):
        for field, src, col, fmt in spec:
            cell = ws.cell(row=8 + i, column=col)
            assert cell.value == pytest.approx(src[period][field], rel=1e-15, abs=0), (period, field)
            assert cell.number_format == fmt, (period, field)


def test_4_the_template_header_and_widths_are_untouched(profiles, tmp_path):
    path, _ = _build(profiles, tmp_path)
    ws = openpyxl.load_workbook(path)["EV Summary"]
    tws = openpyxl.load_workbook(TEMPLATE)["EV Summary"]
    assert [c.value for c in ws[7]] == [c.value for c in tws[7]] == list(gen.HEADERS)
    for letter, width in zip("ABCDEFG", gen.WIDTHS):
        assert ws.column_dimensions[letter].width == tws.column_dimensions[letter].width == width


def test_4_methods_list_the_sources_and_the_manifest_is_hidden_and_parses(profiles, tmp_path):
    path, env = _build(profiles, tmp_path)
    wb = openpyxl.load_workbook(path)
    rows = {r[0].value: [c.value for c in r] for r in wb["Methods"].iter_rows(min_row=2)}
    mod = REPO / "agent_fleet/finance_agent/measure_modules/tcpi.py"
    assert rows["fin_tcpi"][1] == "fin_tcpi"
    assert rows["fin_tcpi"][2] == engine.DECLARED["fin_tcpi"].name
    assert rows["fin_tcpi"][4] == hashlib.sha256(mod.read_bytes()).hexdigest()
    assert set(rows) == {"fin_performance_indices", "fin_tcpi"}
    mf = wb["_manifest"]
    assert mf.sheet_state == "hidden"
    assert json.loads(mf["A1"].value) == env
    assert env["kind"] == "cdrl" and env["program_id"] == PROGRAM


# -- arm 5: reproducibility --------------------------------------------------------------------

def test_5_two_builds_agree_and_a_changed_fact_moves_the_hash(profiles, tmp_path):
    _, a = _build(profiles, tmp_path / "a")
    _, b = _build(profiles, tmp_path / "b")
    assert a["cells_sha256"] == b["cells_sha256"]
    assert [s["rows_sha256"] for s in a["sources"]] == [s["rows_sha256"] for s in b["sources"]]
    state = copy.deepcopy(engine.STATE)
    state.facts = [replace(state.facts[0], bcwp=state.facts[0].bcwp + 12345.0), *state.facts[1:]]
    engine.STATE, saved = state, engine.STATE
    try:
        _, c = _build(profiles, tmp_path / "c", state=state)
    finally:
        engine.STATE = saved
    assert c["cells_sha256"] != a["cells_sha256"]
    assert [s["rows_sha256"] for s in c["sources"]] != [s["rows_sha256"] for s in a["sources"]]


def test_5_the_hash_is_over_the_cells_written_and_the_manifest_is_not_one_of_them(profiles, tmp_path):
    path, env = _build(profiles, tmp_path)
    wb = openpyxl.load_workbook(path)
    written = []
    for ws in wb.worksheets:
        if ws.title == "_manifest":
            continue
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                if ws.title == "EV Summary" and (c.row == 7 or c.column == 1 and c.row <= 5
                                                 or c.coordinate == "A1"):
                    continue  # the template's own text, not written by the producer
                written.append([ws.title, c.coordinate, c.value])
    want = hashlib.sha256(json.dumps(
        sorted(written, key=lambda c: (c[0], c[1])), sort_keys=True, separators=(",", ":"),
        default=str).encode()).hexdigest()
    assert env["cells_sha256"] == want


# -- arm 6: undefined --------------------------------------------------------------------------

def test_6_a_spent_budget_writes_the_reason_not_a_number(profiles, tmp_path, monkeypatch):
    state = copy.deepcopy(engine.STATE)
    program = next(p for p in state.programs if p.program_id == PROGRAM)
    ca = {c.ca_id for c in state.accounts_of(PROGRAM)}
    wps = {w.wp_id for w in state.work_packages if w.ca_id in ca}
    state.facts = [replace(f, acwp=f.acwp + 10 * program.bac) if f.wp_id in wps else f
                   for f in state.facts]
    monkeypatch.setattr(engine, "STATE", state)
    path, _ = _build(profiles, tmp_path, state=state)
    ws = openpyxl.load_workbook(path)["EV Summary"]
    cells = [ws.cell(row=r, column=7).value for r in range(8, 40) if ws.cell(row=r, column=1).value]
    assert cells and all(v == "undefined: budget_spent" for v in cells), cells


# -- arm 7: engine routes ----------------------------------------------------------------------

class Double:
    def __init__(self):
        self.calls = []

    def __call__(self, program, caller_id, **kw):
        self.calls.append((program, caller_id))
        return (program, caller_id) == (PROGRAM, ALICE)


@pytest.fixture
def served(profiles, monkeypatch, tmp_path):
    d = Double()
    monkeypatch.setattr(engine, "can_view_program", d)
    root = tmp_path / "root"
    root.mkdir()
    monkeypatch.setattr(engine, "_repo_root", lambda: root)
    with TestClient(engine.app) as c:
        yield c, root / "dist", d


def _post(c, profile=PROFILE, user=ALICE):
    return c.post("/cdrl_export", json={"profile": profile, "program_id": PROGRAM, "user_email": user})


def test_7_profiles_are_header_keys_only(served):
    c, _, _ = served
    r = c.get("/cdrl_profiles")
    assert r.status_code == 200
    got = r.json()
    assert PROFILE in [x["name"] for x in got]
    assert all(set(x) == {"name", "format", "scope", "label"} for x in got)


def test_7_an_unknown_profile_is_404_unknown_profile(served):
    c, _, _ = served
    r = _post(c, profile="nope")
    assert r.status_code == 404
    assert r.json()["detail"]["error"] == "unknown_profile"


def test_7_a_non_member_is_refused_byte_equal_with_no_file_and_no_source_measured(served, monkeypatch):
    c, dist, d = served
    built, ran = [], []
    monkeypatch.setattr(cdrl, "build_cdrl", lambda *a, **k: built.append(1))
    monkeypatch.setattr(engine, "_measure_envelope", lambda *a, **k: ran.append(1))
    r = _post(c, user=BOB)
    assert r.status_code == 403
    assert r.json()["detail"] == ENVELOPE
    assert not built and not ran, "the gate ran after the build"
    assert not dist.exists() or not list(dist.iterdir())
    assert d.calls == [(PROGRAM, BOB)]


def test_7_a_member_gets_the_locator_and_the_envelope(served):
    c, dist, _ = served
    r = _post(c)
    assert r.status_code == 200, r.text
    b = r.json()
    f = dist / b["artifact_filename"]
    assert f.is_file()
    assert b["artifact_sha256"] == "sha256:" + hashlib.sha256(f.read_bytes()).hexdigest()
    assert b["artifact_bytes"] == f.stat().st_size
    assert b["envelope"]["kind"] == "cdrl" and b["envelope"]["cells_sha256"]
    assert json.loads(openpyxl.load_workbook(f)["_manifest"]["A1"].value) == b["envelope"]


def test_7_the_artifact_route_serves_the_file_to_a_member_only(served, monkeypatch):
    c, dist, _ = served
    name = _post(c).json()["artifact_filename"]
    assert c.get(f"/artifact/{name}").status_code == 403          # no bearer: not a member
    monkeypatch.setattr(engine, "_current_caller", lambda: SimpleNamespace(authz_id=ALICE))
    r = c.get(f"/artifact/{name}")
    assert r.status_code == 200
    assert r.content == (dist / name).read_bytes()
    assert c.get("/artifact/fin-cdrl-nope-NOPE.xlsx").status_code == 404


# --- the flattened image: no checkout, no package builder, no Pyodide runtime -------------------
# The finance image ships neither `scripts/build_fin_package.py` nor `.pyodide-cache`, so
# `_repo_root()` is None in the pod. A CDRL is an xlsx and needs neither: it is written to
# `_cdrl_out_dir()`, which BOTH `/cdrl_export` and `/artifact` read.

@pytest.fixture
def pod(profiles, monkeypatch, tmp_path):
    """The deployed pod: `_repo_root` finds nothing."""
    d = Double()
    monkeypatch.setattr(engine, "can_view_program", d)
    monkeypatch.setattr(engine, "_repo_root", lambda: None)
    monkeypatch.delenv("FIN_CDRL_ARTIFACT_DIR", raising=False)
    with TestClient(engine.app) as c:
        yield c, d


def test_9_without_a_checkout_the_env_dir_is_where_a_cdrl_is_written_and_served(pod, monkeypatch, tmp_path):
    c, _ = pod
    out = tmp_path / "cdrl-out"                       # does not exist yet: created on first write
    monkeypatch.setenv("FIN_CDRL_ARTIFACT_DIR", str(out))
    r = _post(c)
    assert r.status_code == 200, r.text
    name = r.json()["artifact_filename"]
    assert (out / name).is_file()
    assert c.get(f"/artifact/{name}").status_code == 403          # no bearer: not a member
    monkeypatch.setattr(engine, "_current_caller", lambda: SimpleNamespace(authz_id=ALICE))
    got = c.get(f"/artifact/{name}")
    assert got.status_code == 200 and got.content == (out / name).read_bytes()


def test_9_with_no_env_the_functions_own_default_is_used(pod, monkeypatch, tmp_path):
    c, _ = pod
    fallback = tmp_path / "workspace" / "fin_cdrl"
    monkeypatch.setattr(engine, "_CDRL_FALLBACK_DIR", fallback)   # the default, NOT the env
    assert engine._cdrl_out_dir() == fallback
    r = _post(c)
    assert r.status_code == 200, r.text
    name = r.json()["artifact_filename"]
    assert (fallback / name).is_file()
    monkeypatch.setattr(engine, "_current_caller", lambda: SimpleNamespace(authz_id=ALICE))
    assert c.get(f"/artifact/{name}").status_code == 200


def test_9_in_a_checkout_the_dist_dir_is_unchanged(served, monkeypatch):
    _, dist, _ = served
    monkeypatch.delenv("FIN_CDRL_ARTIFACT_DIR", raising=False)
    assert engine._cdrl_out_dir() == dist


def test_9_the_503_is_only_a_directory_that_cannot_be_written(pod, monkeypatch, tmp_path):
    c, _ = pod
    blocker = tmp_path / "a-file"
    blocker.write_text("not a directory")
    monkeypatch.setenv("FIN_CDRL_ARTIFACT_DIR", str(blocker / "sub"))
    r = _post(c)
    assert r.status_code == 503
    assert "cannot be created or written" in r.json()["detail"]


# --- who may fetch a CDRL artifact: the program's members, by the caller's own bearer ------------

def test_10_a_non_member_bearer_and_a_service_bearer_are_refused_a_member_is_served(served, monkeypatch):
    c, dist, d = served
    name = _post(c).json()["artifact_filename"]
    for who, expect in ((BOB, 403), ("svc:supervisor", 403), (None, 403), (ALICE, 200)):
        monkeypatch.setattr(
            engine, "_current_caller",
            (lambda: None) if who is None else (lambda w=who: SimpleNamespace(authz_id=w)),
        )
        r = c.get(f"/artifact/{name}")
        assert r.status_code == expect, (who, r.status_code)
        if expect == 403:
            assert r.json()["detail"] == ENVELOPE, who      # and it never served the bytes
